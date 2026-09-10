"""Parseo de Excel de tickets con openpyxl.

Normaliza las columnas al mismo esquema que normalize() del prototipo.
Acepta encabezados con variaciones comunes (español/inglés, con/sin espacios).
"""

import io

from openpyxl import load_workbook

# Mapa: campo normalizado -> lista de encabezados aceptados (en minúsculas, sin espacios extra)
COLUMN_ALIASES: dict[str, list[str]] = {
    "ticket_id": ["ticket id", "ticketid", "ticket_id", "id"],
    "ticket_type": ["ticket type", "tickettype", "type"],
    "priority": ["priority", "prioridad"],
    "status": ["status", "estado"],
    "queue_status": ["queue status", "queuestatus"],
    "waiting_reason": ["waiting reason", "waitingreason"],
    "subject": ["subject", "asunto"],
    "processor": ["processor", "procesador"],
    "processing_queue": ["processing queue", "processingqueue"],
    "system_role": ["system role", "systemrole"],
    "service_area": ["service area", "servicearea"],
    "customer_name": ["customer name", "customername", "customer"],
    "reported_at": ["reported at", "reportedat"],
    "metric": ["metric"],
    "metric_due_at": ["metric due at", "metricdueat"],
    "cycle_due_at": ["cycle due at", "cycledueat"],
}


def _norm_header(h) -> str:
    return str(h or "").strip().lower()


def parse_tickets(content: bytes) -> list[dict]:
    """Devuelve una lista de dicts normalizados (solo filas con ticket_id)."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []

    # Construye índice columna->campo normalizado
    header_norm = [_norm_header(h) for h in header]
    col_field: dict[int, str] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for idx, h in enumerate(header_norm):
            if h in aliases:
                col_field[idx] = field
                break

    tickets: list[dict] = []
    for raw in rows:
        record = {f: "" for f in COLUMN_ALIASES}
        for idx, value in enumerate(raw):
            field = col_field.get(idx)
            if field:
                record[field] = str(value).strip() if value is not None else ""
        if record["ticket_id"]:
            tickets.append(record)

    wb.close()
    return tickets
